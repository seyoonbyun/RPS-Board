# -*- coding: utf-8 -*-
r"""RPS Board 어드민의 **챕터 삭제**를 감지해 RPI 집계 시트에서 그 행을 지운다.

    python rpi_watch.py            # dry-run (무엇을 지울지만 찍는다)
    python rpi_watch.py --apply    # 실제 제거 + 이메일 보고
    python rpi_watch.py --replay 5 # 커서를 무시하고 최근 5건을 다시 본다 (점검용)

왜 필요한가
    RPS Board 는 지역 담당자가 실시간으로 관리하지만 `RPI : BNI K. All` 은 수동이라,
    챕터가 문을 닫아도 행이 남는다. 멤버가 `_source` 에서 빠지면 분모가 0 이 되고
    수식의 `IFERROR` 가 0 을 돌려주므로 **폐쇄 챕터가 RPI 0.00 짜리 활동 챕터처럼 보인다.**

감지원
    모 시트 `ChapterLog` 탭. 어드민 동작이 한 줄씩 남는다.
        Timestamp | Admin Email | Action | Details
        2026-04-15,01-08-14 | joy.byun@… | 챕터 삭제 | 챕터: 썬

⚠ **로그만 믿고 지우지 않는다.** 담당자가 실수로 지웠다 되살리는 일이 실제로 있다
   (`썬` 은 생성 8초 뒤 삭제됐다). 로그와 `_source`(멤버 0명) 가 **둘 다** 폐쇄를
   가리킬 때만 제거한다.

⚠ **로그의 챕터명은 한글일 수도 영문일 수도 있다.** 담당자가 입력한 그대로 남는다
   (`Unique`·`ONLY ONE` vs `썬`·`에피소드`). RPI 시트는 영문명 기준이라,
   챕터 마스터의 영문·한글을 모두 색인해 해석한다.
   **해석하지 못하면 지우지 않고 보고만 한다** — 조용히 넘기면 폐쇄가 영영 반영되지 않는다.
"""
from __future__ import annotations

import argparse
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")
sys.path.insert(0, str(Path(__file__).resolve().parent))

import paths                                              # noqa: E402
from google_auth import user_credentials                  # noqa: E402
from googleapiclient.discovery import build                # noqa: E402
from rpi_sheet import SID, find_master, read_all_status    # noqa: E402

MASTER_ID = "1JM37uOEu64D0r6zzKggOsA9ZdcK4wBCx0rpuNoVcIYg"   # MY PowerTeam (archive)
LOG_TAB = "ChapterLog"
ACTION_DELETE = "챕터 삭제"


def state_path() -> Path:
    d = paths.snap_dir() if hasattr(paths, "snap_dir") else Path.cwd()
    d.mkdir(parents=True, exist_ok=True)
    return d / "rpi_watch_state.json"


def load_state() -> dict | None:
    """없으면 None. 최초 실행은 호출부가 커서를 로그 끝으로 맞추고 아무것도 하지 않는다."""
    p = state_path()
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def name_index(master_path: Path) -> dict[str, str]:
    """{영문·한글 모두(casefold): 영문명}. 로그에 남은 이름을 영문명으로 해석한다."""
    wb = openpyxl.load_workbook(master_path, data_only=True, read_only=True)
    idx: dict[str, str] = {}
    for sheet in ("ALL 챕터",):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        h = {v: i for i, v in enumerate(rows[0]) if v}
        for r in rows[1:]:
            eng = str(r[h["챕터명(Eng)"]] or "").strip()
            kor = str(r[h["챕터명(Kor)"]] or "").strip()
            if not eng:
                continue
            for k in (eng, kor):
                if k:
                    idx.setdefault(k.casefold(), eng)
    return idx


def main() -> int:
    ap = argparse.ArgumentParser(description="어드민 챕터 삭제 감지 → RPI 시트 행 제거")
    ap.add_argument("--apply", action="store_true", help="실제 제거 + 이메일 보고")
    ap.add_argument("--replay", type=int, default=0,
                    help="커서를 무시하고 마지막 N건을 다시 본다 (점검용, 커서는 안 옮긴다)")
    a = ap.parse_args()

    creds = user_credentials()
    svc = build("sheets", "v4", credentials=creds, cache_discovery=False)

    # ── 1) ChapterLog 읽기 ───────────────────────────────────
    log = svc.spreadsheets().values().get(
        spreadsheetId=MASTER_ID, range=f"'{LOG_TAB}'!A2:D100000",
        valueRenderOption="FORMATTED_VALUE").execute().get("values", [])
    st = load_state()
    if st is None and not a.replay:
        # 최초 실행 — 과거 로그는 처리하지 않는다. 소급 정리는 사람이 한 번 하고 시작한다.
        # (그러지 않으면 이미 끝난 옛 삭제 건들이 한꺼번에 보고돼 노이즈가 된다.)
        state_path().write_text(json.dumps(
            {"cursor": len(log),
             "note": "최초 실행 — 기존 로그는 처리하지 않았다",
             "checked_at": datetime.datetime.now().isoformat(timespec="seconds")},
            ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"최초 실행 — 커서를 {len(log)} 로 맞추고 종료한다 (과거 로그는 처리하지 않는다).")
        print(f"   상태 파일 {state_path()}")
        return 0
    st = st or {}
    cursor = 0 if a.replay else int(st.get("cursor", 0))
    fresh = log[-a.replay:] if a.replay else log[cursor:]
    print(f"{LOG_TAB} {len(log)}행 · 커서 {cursor} · 신규 {len(fresh)}행")

    deletions = []
    for row in fresh:
        row = list(row) + [""] * (4 - len(row))
        ts, who, action, details = (str(x).strip() for x in row[:4])
        if action != ACTION_DELETE:
            continue
        m = re.search(r"챕터\s*:\s*(.+?)\s*$", details)
        deletions.append({"ts": ts, "who": who, "raw": m.group(1) if m else details})

    if not deletions:
        print("삭제 로그 없음 — 할 일 없음")
        if a.apply and not a.replay:
            st["cursor"] = len(log)
            st["checked_at"] = datetime.datetime.now().isoformat(timespec="seconds")
            state_path().write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
        return 0

    # ── 2) 이름 해석 + 교차검증 ───────────────────────────────
    master_path = find_master()
    idx = name_index(master_path)
    status = read_all_status(master_path)

    src = svc.spreadsheets().values().get(
        spreadsheetId=SID, range="'_source'!B2:B5000",
        valueRenderOption="UNFORMATTED_VALUE").execute().get("values", [])
    src_count: dict[str, int] = {}
    for r in src:
        if r and str(r[0]).strip():
            k = str(r[0]).strip().casefold()
            src_count[k] = src_count.get(k, 0) + 1

    ch_rows = svc.spreadsheets().values().get(
        spreadsheetId=SID, range="'챕터 RPI'!A2:G5000",
        valueRenderOption="FORMULA").execute().get("values", [])
    ch_rows = [list(r) + [""] * (7 - len(r)) for r in ch_rows if len(r) > 2 and str(r[2]).strip()]
    present = {str(r[2]).strip().casefold(): str(r[2]).strip() for r in ch_rows[1:]}

    todo, skipped = [], []
    for d in deletions:
        eng = idx.get(d["raw"].casefold())
        if not eng:
            skipped.append({**d, "why": "챕터명을 영문명으로 해석하지 못했다 (마스터에 없는 이름)"})
            continue
        d["eng"] = eng
        if eng.casefold() not in present:
            skipped.append({**d, "why": f"RPI 시트에 {eng} 행이 이미 없다"})
            continue
        n = src_count.get(eng.casefold(), 0)
        if n > 0:
            skipped.append({**d, "why": f"_source 에 아직 {n}명 남아 있다 (되살렸거나 반영 전)"})
            continue
        todo.append(d)

    print(f"\n삭제 로그 {len(deletions)}건 → 제거 대상 {len(todo)}건 · 보류 {len(skipped)}건")
    for d in todo:
        stt = status.get(d["eng"].casefold(), ("?", "?"))
        print(f"   - {d['eng']:<14} {stt[1]:<18} {d['ts']}  by {d['who']}")
    for d in skipped:
        print(f"   · 보류 {d['raw']:<14} {d['why']}")

    if not a.apply:
        print("\n[dry-run] 쓰지 않았다. --apply 를 붙일 것.")
        return 0

    body_lines = []
    if todo:
        # 스냅샷 — 되돌리기는 이 파일로 한다(수식까지 들어 있다).
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        snap = paths.snap_dir() / f"rpi_watch_{stamp}.json"
        snap.write_text(json.dumps({"챕터 RPI": ch_rows}, ensure_ascii=False, indent=1),
                        encoding="utf-8")

        drop = {d["eng"].casefold() for d in todo}
        head, keep = ch_rows[0], [r for r in ch_rows[1:] if str(r[2]).strip().casefold() not in drop]
        keep.sort(key=lambda r: str(r[2]).strip().casefold())
        for i, r in enumerate(keep, start=2):
            r[0] = i
        head[0] = 1
        out = [head] + keep
        # ⚠ 행이 줄면 쓰기 범위도 줄어 이전 마지막 행이 남는다 → 빈 행으로 덮는다.
        write = out + [[""] * 7] * max(0, len(ch_rows) - len(out))
        svc.spreadsheets().values().update(
            spreadsheetId=SID, range=f"'챕터 RPI'!A2:G{1 + len(write)}",
            valueInputOption="USER_ENTERED", body={"values": write}).execute()
        print(f"\n✓ 챕터 RPI {len(ch_rows)}행 → {len(out)}행")

        body_lines.append(f"제거 {len(todo)}건")
        for d in todo:
            stt = status.get(d["eng"].casefold(), ("?", "?"))
            body_lines.append(f"  - {d['eng']} ({stt[1]}) · 마스터 상태 {stt[0]}")
            body_lines.append(f"    어드민 삭제 {d['ts']} by {d['who']} · _source 멤버 0명 확인")
        body_lines.append(f"  챕터 RPI {len(ch_rows)}행 → {len(out)}행")
        body_lines.append(f"  스냅샷 {snap.name}")

    if skipped:
        body_lines.append(f"\n보류 {len(skipped)}건 (사람 확인 필요)")
        for d in skipped:
            body_lines.append(f"  - {d['raw']} ({d['ts']} by {d['who']}) — {d['why']}")

    # 아무 일도 없었으면 메일을 보내지 않는다. 매시간 빈 메일이 쌓이면 아무도 안 본다.
    if body_lines:
        from notify import send_email
        subject = f"[RPI] 챕터 삭제 반영 {len(todo)}건" + (f" · 보류 {len(skipped)}건" if skipped else "")
        try:
            send_email(subject, "\n".join(body_lines))
            print("✓ 보고 메일 발송")
        except Exception as e:                                    # noqa: BLE001
            print(f"⚠ 메일 발송 실패 (제거는 완료됨): {e}")

    if not a.replay:
        st["cursor"] = len(log)
        st["checked_at"] = datetime.datetime.now().isoformat(timespec="seconds")
        state_path().write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
